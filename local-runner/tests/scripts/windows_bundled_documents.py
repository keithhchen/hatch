"""Real Windows Runner fixture driver. Not visual/Desktop UAT; never skip dependencies.

Invoked only by windows_execution.rs using the Runner-configured bundled Python.
Uses shipped Skill render commands rather than implementing another renderer.
"""
import hashlib
import importlib
import json
import os
from pathlib import Path
import subprocess
import sys


def require(condition, message):
    if not condition:
        raise RuntimeError(message)


def bundled(name, root, directory=False):
    value = os.environ.get(name)
    require(value, "Runner did not configure " + name)
    result = Path(value).resolve(strict=True)
    require(result.is_relative_to(root), name + " escapes bundled runtime: " + str(result))
    require(result.is_dir() if directory else result.is_file(), name + " has wrong file type")
    return result


def run(argv, label):
    result = subprocess.run([str(arg) for arg in argv], capture_output=True,
                            encoding="utf-8", errors="replace", timeout=100)
    Path(label + ".stdout.txt").write_text(result.stdout, encoding="utf-8")
    Path(label + ".stderr.txt").write_text(result.stderr, encoding="utf-8")
    require(result.returncode == 0,
            f"{label} exited {result.returncode}: {result.stdout}\n{result.stderr}")
    return result


def check_png(file):
    from PIL import Image, ImageChops
    with Image.open(file) as image:
        image.load()
        require(image.format == "PNG" and min(image.size) >= 100, "invalid render: " + str(file))
        gray = image.convert("L")
        require(ImageChops.difference(gray, Image.new("L", gray.size, 255)).getbbox(),
                "blank render: " + str(file))
    return {"file": str(file), "bytes": file.stat().st_size,
            "sha256": hashlib.sha256(file.read_bytes()).hexdigest()}


def generate(node):
    from docx import Document
    from openpyxl import Workbook, load_workbook
    from pptx import Presentation
    marker = "自动化测试 非UAT Hatch 314159"
    document = Document()
    document.add_heading(marker, 0)
    document.add_paragraph("中文文档生成与渲染检查")
    document.save("测试 文档.docx")
    workbook = Workbook()
    workbook.active["A1"] = marker
    workbook.active["A2"] = 21
    workbook.active["B2"] = "=A2*2"
    workbook.active.column_dimensions["A"].width = 48
    workbook.save("测试 表格.xlsx")
    # Node must actually generate the PPTX, not just print a version.
    code = r"""
const fs = require('node:fs');
const path = require('node:path');
const nodePath = process.env.NODE_PATH;
if (nodePath !== process.env.HATCH_NODE_MODULES) throw Error('Runner module paths disagree');
if (nodePath.startsWith('\\\\?\\')) throw Error('Runner leaked a verbatim NODE_PATH');
const modulePath = require.resolve('pptxgenjs');
const root = fs.realpathSync(process.env.HATCH_NODE_MODULES);
const relative = path.relative(root, fs.realpathSync(modulePath));
if (relative.startsWith('..') || path.isAbsolute(relative)) throw Error('Non-bundled pptxgenjs');
const PptxGenJS = require('pptxgenjs');
const pptx = new PptxGenJS();
pptx.addSlide().addText(process.argv[2], {x:1, y:1, w:8, h:2, fontSize:28});
pptx.writeFile({fileName:process.argv[1]}).then(() => {
  console.log(JSON.stringify({executable:process.execPath, module:modulePath, nodePath}));
}).catch(error => {console.error(error); process.exitCode=1;});
"""
    output = run([node, "-e", code, "测试 演示.pptx", marker], "node-generate")
    require(marker in "\n".join(p.text for p in Document("测试 文档.docx").paragraphs), "DOCX lost text")
    restored = load_workbook("测试 表格.xlsx")
    require(restored.active["A1"].value == marker and restored.active["B2"].value == "=A2*2", "XLSX lost cells")
    restored.close()
    restored_slides = Presentation("测试 演示.pptx")
    require(marker in "\n".join(shape.text for slide in restored_slides.slides
                               for shape in slide.shapes if shape.has_text_frame), "PPTX lost text")
    return {"node": json.loads(output.stdout), "files": ["测试 文档.docx", "测试 演示.pptx", "测试 表格.xlsx"]}


def render(stage, skills):
    from pypdf import PdfReader
    script, filename = {
        "docx": ("documents/scripts/render_docx.py", "测试 文档.docx"),
        "pptx": ("presentations/scripts/pptx_tool.py", "测试 演示.pptx"),
        "xlsx": ("spreadsheets/scripts/xlsx_tool.py", "测试 表格.xlsx"),
    }[stage]
    command = [sys.executable, "-X", "utf8", skills / script]
    if stage != "docx":
        command.append("render")
    command += [filename, "--output-dir", stage + "-render"]
    result = json.loads(run(command, stage).stdout)
    require(result.get("status") == "ok" and not result.get("warning"), "render did not fully succeed: " + str(result))
    pdf = Path(result["pdf"])
    require(pdf.is_file() and len(PdfReader(pdf).pages) > 0, "missing/empty PDF")
    pages = result.get("pages", [])
    require(pages, "missing PNG previews; PDF-only fallback is not a pass")
    return {"pdf": str(pdf), "pages": [check_png(Path(file)) for file in pages]}


def cjk(native, pdftoppm, pdfinfo):
    from PIL import Image
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from pypdf import PdfReader
    prefix = native / "poppler/Library"
    for relative in ["share/poppler/cMap/Adobe-GB1/UniGB-UCS2-H", "share/poppler/cidToUnicode/Adobe-GB1"]:
        require((prefix / relative).is_file(), "missing bundled Chinese mapping: " + relative)
    config = Path(os.environ["FONTCONFIG_FILE"]).resolve(strict=True)
    require(config == (prefix / "etc/fonts/fonts.conf").resolve(strict=True), "wrong Fontconfig config")
    require(Path(os.environ["FONTCONFIG_PATH"]).resolve(strict=True) == config.parent, "wrong Fontconfig directory")
    fonts = run([prefix / "bin/fc-list.exe", ":lang=zh", "family"], "chinese-fonts")
    require(fonts.stdout.strip() and not fonts.stderr.strip(), "Fontconfig cannot discover Chinese fonts")
    output = Path("cjk-render")
    output.mkdir()
    pdf = output / "非嵌入 中文.pdf"
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    page = canvas.Canvas(str(pdf), pagesize=(360, 160), pageCompression=0)
    page.setFont("STSong-Light", 28)
    for index, character in enumerate("中文字形测试"):
        page.drawString(36 + index * 40, 80, character)
    page.save()
    reader = PdfReader(pdf)
    cid_fonts = [font.get_object() for font in reader.pages[0]["/Resources"]["/Font"].values()
                 if font.get_object().get("/Subtype") == "/Type0"]
    require(cid_fonts and str(cid_fonts[0]["/Encoding"]) == "/UniGB-UCS2-H", "fixture does not exercise named CMap")
    for font in cid_fonts:
        for child in font["/DescendantFonts"]:
            descriptor = child.get_object()["/FontDescriptor"].get_object()
            require(not any(key in descriptor for key in ["/FontFile", "/FontFile2", "/FontFile3"]), "fixture font was embedded")
    run([pdfinfo, pdf], "chinese-pdfinfo")
    rendered = run([pdftoppm, "-f", "1", "-singlefile", "-r", "72", "-png", pdf, output / "chinese"], "chinese-render")
    require(not rendered.stderr.strip(), "Poppler CMap/font diagnostics: " + rendered.stderr)
    png = output / "chinese.png"
    report = check_png(png)
    # Reject blank output and repeated tofu boxes; this is not OCR/visual UAT.
    with Image.open(png) as image:
        fingerprints = set()
        for index in range(6):
            crop = image.convert("L").crop((36 + index * 40, 48, 76 + index * 40, 88))
            mask = crop.point(lambda value: 0 if value < 180 else 255)
            require(mask.getextrema()[0] == 0, "missing Chinese glyph")
            fingerprints.add(hashlib.sha256(mask.tobytes()).hexdigest())
        require(len(fingerprints) >= 4, "repeated missing-glyph boxes instead of Chinese")
    return {"pdf": str(pdf), "png": report, "fontconfig": str(config), "fonts": fonts.stdout.strip()}


def main():
    require(sys.platform == "win32", "Windows-only integration; other OS cannot prove Windows")
    root = Path(os.environ["HATCH_RUNTIME_ROOT"]).resolve(strict=True)
    manifest = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    require(manifest["target"]["platform"] == "win32", "not a Windows bundle")
    python = bundled("HATCH_PYTHON", root)
    require(Path(sys.executable).resolve() == python, "not running bundled Python")
    node = bundled("HATCH_NODE", root)
    native = bundled("HATCH_NATIVE_RUNTIME_ROOT", root, True)
    skills = bundled("HATCH_DOCUMENT_SKILLS_ROOT", root, True)
    soffice = bundled("HATCH_SOFFICE", root)
    pdftoppm = bundled("HATCH_PDFTOPPM", root)
    pdfinfo = bundled("HATCH_PDFINFO", root)
    bundled("HATCH_NODE_MODULES", root, True)
    for executable, relative in [(python, manifest["python"]["executable"]),
                                 (node, manifest["node"]["executable"]),
                                 (soffice, manifest["native"]["binaries"]["soffice"]),
                                 (pdftoppm, manifest["native"]["binaries"]["pdftoppm"]),
                                 (pdfinfo, manifest["native"]["binaries"]["pdfinfo"])]:
        require(executable == (root / relative).resolve(strict=True), "executable disagrees with bundle manifest")
    packages = (root / manifest["python"]["package_root"]).resolve(strict=True)
    require(packages.is_relative_to(root), "Python packages escape runtime")
    for name in ["docx", "openpyxl", "pptx", "PIL", "reportlab", "pypdf"]:
        module = importlib.import_module(name)
        require(Path(module.__file__).resolve().is_relative_to(packages), "non-bundled Python dependency: " + name)
    stage = sys.argv[1]
    result = generate(node) if stage == "generate" else cjk(native, pdftoppm, pdfinfo) if stage == "cjk" else render(stage, skills)
    report = {"status": "ok", "fixture_not_uat": True, "stage": stage, "runtime": str(root),
              "manifest_sha256": hashlib.sha256((root / "manifest.json").read_bytes()).hexdigest(), **result}
    Path("report-" + stage + ".json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
