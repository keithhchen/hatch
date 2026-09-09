"""Opt-in real bundled toolchain coverage, not installed Desktop UAT.

Run with the bundled Python and HATCH_RUNTIME_ROOT pointing at its runtime.
All generated documents are test-only and confined to TemporaryDirectory.
"""
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

SKILLS = Path(__file__).resolve().parents[1] / "skills"
sys.path.insert(0, str(SKILLS))
from _shared.libreoffice import run_libreoffice


@unittest.skipUnless(os.environ.get("HATCH_RUNTIME_ROOT"), "requires bundled runtime")
class LegacyOfficeIntegration(unittest.TestCase):
    def test_legacy_round_trips(self):
        from docx import Document
        from openpyxl import Workbook, load_workbook
        from pptx import Presentation

        runtime = Path(os.environ["HATCH_RUNTIME_ROOT"]).resolve()
        manifest = json.loads((runtime / "manifest.json").read_text())
        soffice = runtime / manifest["native"]["binaries"]["soffice"]
        environment = dict(os.environ, HATCH_SOFFICE=str(soffice))
        marker = "Hatch legacy roundtrip 314159"
        with tempfile.TemporaryDirectory(prefix="hatch-legacy-office-") as temporary:
            root = Path(temporary)
            source = root / "source"
            source.mkdir()
            document = Document()
            document.add_paragraph(marker)
            document.save(source / "document.docx")
            workbook = Workbook()
            workbook.active["A1"] = marker
            workbook.active["B1"] = 42
            workbook.save(source / "workbook.xlsx")
            presentation = Presentation()
            presentation.slides.add_slide(presentation.slide_layouts[0]).shapes.title.text = marker
            presentation.save(source / "presentation.pptx")

            for filename, legacy, modern in [
                ("document.docx", "doc", "docx"),
                ("document.docx", "rtf", "docx"),
                ("workbook.xlsx", "xls", "xlsx"),
                ("presentation.pptx", "ppt", "pptx"),
            ]:
                with self.subTest(format=legacy):
                    original = source / filename
                    original_digest = hashlib.sha256(original.read_bytes()).digest()
                    legacy_dir = root / legacy
                    legacy_dir.mkdir()
                    profile = root / (legacy + "-profile")
                    profile.mkdir()
                    result = run_libreoffice([
                        str(soffice), "--headless", "--nologo", "--norestore",
                        "-env:UserInstallation=" + profile.as_uri(),
                        "--convert-to", legacy, "--outdir", str(legacy_dir), str(original)
                    ], profile=profile, environment=environment, timeout=120)
                    self.assertEqual(result.returncode, 0, result.stderr)
                    legacy_file = legacy_dir / (original.stem + "." + legacy)
                    self.assertTrue(legacy_file.is_file(), result.stdout)
                    digest = hashlib.sha256(legacy_file.read_bytes()).digest()
                    output = root / (legacy + "-converted")
                    converted = subprocess.run([
                        sys.executable, str(SKILLS / "documents/scripts/office_convert.py"),
                        str(legacy_file), "--format", modern, "--output-dir", str(output)
                    ], env=environment, capture_output=True, text=True, timeout=180)
                    self.assertEqual(converted.returncode, 0, converted.stdout + converted.stderr)
                    self.assertEqual(json.loads(converted.stdout)["status"], "ok")
                    target = output / (original.stem + "." + modern)
                    if modern == "docx":
                        self.assertIn(marker, "\n".join(p.text for p in Document(target).paragraphs))
                    elif modern == "xlsx":
                        restored = load_workbook(target)
                        self.assertEqual(restored.active["A1"].value, marker)
                        self.assertEqual(restored.active["B1"].value, 42)
                        restored.close()
                    else:
                        restored = Presentation(target)
                        self.assertIn(marker, "\n".join(shape.text for slide in restored.slides
                            for shape in slide.shapes if shape.has_text_frame))
                    self.assertEqual(hashlib.sha256(original.read_bytes()).digest(), original_digest)
                    self.assertEqual(hashlib.sha256(legacy_file.read_bytes()).digest(), digest)


if __name__ == "__main__":
    unittest.main(verbosity=2)
