import errno, hashlib, json, os, pathlib, subprocess, sys

root = pathlib.Path(sys.argv[1])
checks = []
def check(name, fn):
    try:
        fn()
        checks.append({'name': name, 'status': 'passed'})
    except Exception as e:
        checks.append({'name': name, 'status': 'failed', 'error': repr(e)})

def denied(fn):
    try:
        fn()
    except OSError as e:
        if e.errno in (errno.EACCES, errno.EPERM) or getattr(e, 'winerror', None) == 5:
            return
        raise
    raise AssertionError('operation was allowed')

def write_workspace():
    p = root / 'workspace' / 'python-marker'
    p.write_text('python-ok')
    assert p.read_text() == 'python-ok'

def environment_clean():
    forbidden = {'HATCH_PROBE_PASSWORD', 'HATCH_PROBE_ENV_CANARY', 'GITHUB_TOKEN', 'GH_TOKEN',
        'AWS_SECRET_ACCESS_KEY', 'OPENAI_API_KEY', 'NODE_OPTIONS'}
    assert not forbidden.intersection(k.upper() for k in os.environ), 'unexpected environment key (values withheld)'

check('environment credentials absent', environment_clean)

check('workspace read/write', write_workspace)
for directory in ('attachments', 'runtime'):
    p = root / directory / 'probe-canary.txt'
    check(directory + ' read', lambda p=p: p.read_text() == 'readonly' or (_ for _ in ()).throw(AssertionError('canary')))
    check(directory + ' write denied', lambda p=p: denied(lambda: p.write_text('ESCAPE')))
check('ungranted read denied', lambda: denied(lambda: (root / 'ungranted' / 'secret.txt').read_text()))
check('ungranted write denied', lambda: denied(lambda: (root / 'ungranted' / 'secret.txt').write_text('ESCAPE')))
check('synthetic internal DB read denied', lambda: denied(lambda: (root / 'ungranted' / 'internal.db').read_bytes()))

def descendant_boundary():
    # Real Python descendant inherits the restricted identity and Job. Never
    # inspect a user's actual DB/credentials, or print canary contents.
    code = '''import errno, os, pathlib, sys
assert "HATCH_PROBE_PASSWORD" not in os.environ
assert "HATCH_PROBE_ENV_CANARY" not in os.environ
for name in ("secret.txt", "internal.db"):
    try: pathlib.Path(sys.argv[1], "ungranted", name).read_bytes()
    except OSError as e:
        if e.errno not in (errno.EACCES, errno.EPERM) and getattr(e, "winerror", None) != 5: raise
    else: raise AssertionError("descendant read allowed")
print("descendant-boundary-ok")
'''
    result = subprocess.run([sys.executable, '-c', code, str(root)], capture_output=True, timeout=20)
    assert result.returncode == 0 and b'descendant-boundary-ok' in result.stdout, 'descendant boundary failed'
check('descendant credentials/internal DB denied and environment clean', descendant_boundary)

# This is the real bundled openpyxl and real bundled soffice, running inside
# the SAME container/job. No host-side conversion or fallback.
def libreoffice():
    import openpyxl
    source = root / 'workspace' / 'formula.xlsx'
    wb = openpyxl.Workbook()
    wb.active['A2'] = 42
    wb.active['A3'] = '=A2+8'
    wb.save(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    out = root / 'workspace' / 'converted'
    out.mkdir()
    exe = sys.argv[2]
    profile = (root / 'scratch' / 'lo-profile').as_uri()
    def convert(src, dest):
        result = subprocess.run([exe, '-env:UserInstallation=' + profile,
            '--headless', '--nologo', '--nodefault', '--nofirststartwizard',
            '--convert-to', 'xlsx', '--outdir', str(dest), str(src)],
            capture_output=True, text=True, errors='replace', timeout=60)
        return {'exit_code': result.returncode, 'stdout': result.stdout[-16384:], 'stderr': result.stderr[-16384:]}
    positive = convert(source, out)
    assert positive['exit_code'] == 0, positive
    result = out / source.name
    assert openpyxl.load_workbook(result, data_only=True).active['A3'].value == 50, positive
    assert openpyxl.load_workbook(result, data_only=False).active['A3'].value == '=A2+8'
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    checks.append({'name': 'LO real recalc cached 50 / formula / source unchanged', 'status': 'passed', 'process': positive})
    # A real valid input produced above is first copied by the HOST coordinator
    # on the second invocation. We do not interpret arbitrary LO failure as an
    # ACL denial: ambiguous diagnostics are explicitly inconclusive.

if len(sys.argv) > 3 and sys.argv[3] == 'lo-negative':
    def negative():
        import openpyxl
        exe = sys.argv[2]
        allowed_output = root / 'workspace' / 'attachment-converted'
        allowed_output.mkdir()
        allowed = subprocess.run([exe, '-env:UserInstallation=' + (root / 'scratch' / 'lo-attachment').as_uri(),
            '--headless', '--convert-to', 'xlsx', '--outdir', str(allowed_output), str(root / 'attachments' / 'allowed.xlsx')],
            capture_output=True, text=True, errors='replace', timeout=60)
        assert allowed.returncode == 0, (allowed.stdout, allowed.stderr)
        assert openpyxl.load_workbook(allowed_output / 'allowed.xlsx', data_only=True).active['A3'].value == 50
        checks.append({'name': 'LO readonly attachment input conversion', 'status': 'passed'})
        for name, src, dest in (
            ('LO ungranted input', root / 'ungranted' / 'hidden.xlsx', root / 'workspace' / 'negative'),
            ('LO readonly attachment output', root / 'workspace' / 'formula.xlsx', root / 'attachments'),
            ('LO readonly runtime output', root / 'workspace' / 'formula.xlsx', root / 'runtime')):
            dest.mkdir(exist_ok=True) if dest.name == 'negative' else None
            result = subprocess.run([exe, '-env:UserInstallation=' + (root / 'scratch' / 'lo-negative').as_uri(),
                '--headless', '--convert-to', 'xlsx', '--outdir', str(dest), str(src)],
                capture_output=True, text=True, errors='replace', timeout=60)
            text = result.stdout + result.stderr
            artifact = dest / src.name
            status = 'failed' if artifact.exists() else 'inconclusive'
            if not artifact.exists() and any(s in text.lower() for s in ('access denied', 'permission denied', 'access is denied')):
                status = 'passed'
            checks.append({'name': name, 'status': status, 'exit_code': result.returncode, 'diagnostics': text[-16384:]})
    check('LO negative invocation', negative)
else:
    check('LO integration', libreoffice)

print(json.dumps({'checks': checks}))
sys.exit(0 if all(c['status'] == 'passed' for c in checks) else 1)
